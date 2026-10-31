"""Small real documents for knowledge tests, built with the same libraries users' tools write them with."""

from __future__ import annotations

import io


def docx(title: str, sections: list[tuple[str, str]]) -> bytes:
    import docx as python_docx

    document = python_docx.Document()
    document.add_heading(title, level=1)
    for heading, body in sections:
        document.add_heading(heading, level=2)
        document.add_paragraph(body)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def xlsx(sheet: str, rows: list[list[object]]) -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    ws = book.active
    assert ws is not None
    ws.title = sheet
    for row in rows:
        ws.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def pptx(slides: list[tuple[str, str]]) -> bytes:
    from pptx import Presentation

    deck = Presentation()
    for title, body in slides:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def pdf(pages: list[str]) -> bytes:
    """A minimal text PDF (Helvetica, one line per page) with a correct xref table."""
    objects: list[bytes] = []
    kids = " ".join(f"{3 + i * 2} 0 R" for i in range(len(pages)))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    font_id = 3 + len(pages) * 2
    for i, line in enumerate(pages):
        content = f"BT /F1 12 Tf 72 720 Td ({line}) Tj ET".encode()
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 {font_id} 0 R >> >> "
            f"/Contents {4 + i * 2} 0 R >>".encode()
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()
