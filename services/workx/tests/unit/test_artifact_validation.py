"""Content validation for agent outputs and uploads (Part 4 §12, Part 13 security tests)."""

from __future__ import annotations

import io
import zipfile

import pytest

from workx.domains.artifacts.validate import Rejected, safe_name, validate


def xlsx() -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.append(["Account", "Revenue"])
    sheet.append(["Northwind Traders", 4200000])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def docx() -> bytes:
    from docx import Document

    document = Document()
    document.add_paragraph("Quarterly summary for Contoso")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def pptx() -> bytes:
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Roadmap"
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def pdf(password: str | None = None) -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    if password:
        writer.encrypt(password)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def zip_with(entries: dict[str, bytes]) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    return out.getvalue()


def test_office_files_are_accepted_with_previews() -> None:
    sheet = validate("accounts.xlsx", xlsx())
    assert sheet.kind == "sheet" and "Northwind Traders,4200000" in sheet.preview
    assert "Contoso" in validate("summary.docx", docx()).preview
    assert "Roadmap" in validate("deck.pptx", pptx()).preview
    assert validate("blank.pdf", pdf()).kind == "pdf"
    assert validate("notes.md", b"# Hello\n").preview == "# Hello\n"


@pytest.mark.parametrize(
    ("name", "data", "reason"),
    [
        (
            "evil.xlsx",
            zip_with({"[Content_Types].xml": b"x", "xl/workbook.xml": b"x", "../../etc/cron.d/x": b"x"}),
            "unsafe_entry_path",
        ),
        (
            "abs.docx",
            zip_with({"[Content_Types].xml": b"x", "word/document.xml": b"x", "/etc/passwd": b"x"}),
            "unsafe_entry_path",
        ),
        (
            "bomb.xlsx",
            zip_with({"[Content_Types].xml": b"x", "xl/workbook.xml": b"\0" * 5_000_000}),
            "compression_ratio",
        ),
        (
            "macro.xlsx",
            zip_with({"[Content_Types].xml": b"x", "xl/workbook.xml": b"x", "xl/vbaProject.bin": b"x"}),
            "macros_not_allowed",
        ),
        ("macro.xlsm", xlsx(), "macros_not_allowed"),
        ("fake.docx", b"not a zip at all", "not_a_valid_office_file"),
        ("wrong.pptx", zip_with({"[Content_Types].xml": b"x", "word/document.xml": b"x"}), "not_a_valid_office_file"),
        ("locked.pdf", pdf("secret"), "encrypted"),
        ("broken.pdf", b"%PDF-1.7 garbage", "not_a_valid_pdf"),
        ("program.md", b"MZ\x90\x00binary", "content_does_not_match_name"),
        ("latin1.txt", "café".encode("latin-1"), "not_utf8_text"),
        ("data.json", b"{nope", "invalid_json"),
        ("tool.exe", b"\x7fELF...", "unsupported_type"),
    ],
)
def test_malicious_or_mislabelled_files_are_rejected(name: str, data: bytes, reason: str) -> None:
    with pytest.raises(Rejected) as rejected:
        validate(name, data)
    assert rejected.value.reason == reason


def test_names_are_reduced_to_a_safe_basename() -> None:
    assert safe_name("../../reports/q3.xlsx") == "q3.xlsx"
    assert safe_name("C:\\\\Users\\\\a\\\\q3.xlsx") == "q3.xlsx"
    for bad in ("", ".env", "a\x00b.txt", "x" * 300):
        with pytest.raises(Rejected):
            safe_name(bad)
